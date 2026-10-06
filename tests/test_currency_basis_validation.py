import ast
from pathlib import Path

from openpyxl import load_workbook

from services.tooling_business_case_calculation import (
    calculate_modelled_tooling_business_case,
)


ROOT = Path(__file__).resolve().parents[1]


def _app_tree():
    return ast.parse((ROOT / "app.py").read_text(encoding="utf-8"))


def _app_source():
    return (ROOT / "app.py").read_text(encoding="utf-8")


def _assigned_values(tree, variable_name):
    return [
        node.value
        for node in ast.walk(tree)
        if isinstance(node, ast.Assign)
        and any(
            isinstance(target, ast.Name) and target.id == variable_name
            for target in node.targets
        )
    ]


def _local_currency_for(region_name):
    workbook_path = ROOT / "Masterdata" / "Regions.xlsx"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"]
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        region_index = headers.index("Region Name")
        currency_index = headers.index("Currency")
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[region_index] == region_name:
                return row[currency_index]
    finally:
        workbook.close()
    raise AssertionError(f"Region {region_name!r} not found in Regions.xlsx")


def test_calculation_currency_is_eur():
    values = _assigned_values(_app_tree(), "calculation_currency")
    assert any(
        isinstance(value, ast.Constant) and value.value == "EUR"
        for value in values
    )


def test_denmark_local_currency_remains_dkk():
    assert _local_currency_for("Denmark") == "DKK"


def test_czech_republic_local_currency_remains_czk():
    assert _local_currency_for("Czech Republic") == "CZK"


def test_poland_local_currency_remains_pln():
    assert _local_currency_for("Poland") == "PLN"


def test_usa_local_currency_remains_usd():
    assert _local_currency_for("USA") == "USD"


def test_labour_rate_display_uses_eur_per_hour():
    assert 'f"Labour Rate: €{float(labour_rate):,.2f}/hr"' in _app_source()


def test_calculation_and_local_region_currencies_are_displayed_separately():
    source = _app_source()
    assert 'f"Calculation Currency: {calculation_currency}"' in source
    assert 'f"Local Region Currency: {currency}"' in source


def test_currency_status_gate_uses_calculation_currency_only():
    tree = _app_tree()
    assignment = next(
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.Assign)
        and any(
            isinstance(target, ast.Name) and target.id == "currency_status"
            for target in node.targets
        )
    )
    assert ast.unparse(assignment.value) == (
        "'verified_eur' if calculation_currency == 'EUR' else 'unavailable'"
    )
    assert "currency" not in {
        node.id
        for node in ast.walk(assignment.value)
        if isinstance(node, ast.Name)
    }


def test_modelled_business_case_runs_for_non_eur_local_regions():
    for region in ("Denmark", "Czech Republic", "Poland", "USA"):
        assert _local_currency_for(region) != "EUR"
        result = calculate_modelled_tooling_business_case(
            2.0, 10_000, 20_000, "verified_eur"
        )
        assert result.annual_modelled_recurring_savings_status == "calculated"
        assert result.simple_payback_years_status == "calculated"


def test_app_does_not_introduce_currency_conversion():
    source = _app_source().lower()
    assert "fx_rate" not in source
    assert "exchange_rate" not in source
    assert "currency_conversion" not in source
    assert "convert_currency" not in source


def test_material_labour_and_machine_formulas_remain_unchanged():
    tree = _app_tree()
    assignments = {
        name: [ast.unparse(value) for value in _assigned_values(tree, name)]
        for name in (
            "calculated_material_cost",
            "calculated_labour_cost",
            "calculated_machine_cost",
        )
    }
    assert any(
        value == "float(part_weight) * float(material_cost)"
        for value in assignments["calculated_material_cost"]
    )
    assert any(
        value == "float(cycle_time_sec) / 3600 * float(labour_rate)"
        for value in assignments["calculated_labour_cost"]
    )
    machine_calls = [
        ast.unparse(node)
        for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Name)
        and node.func.id == "calculate_machine_cost"
    ]
    assert "calculate_machine_cost(cycle_time_sec, machine_rate_eur_hr)" in (
        machine_calls
    )


def test_total_should_cost_formula_remains_unchanged():
    values = _assigned_values(_app_tree(), "total_should_cost")
    assert "sum(cost_breakdown.values())" in [ast.unparse(value) for value in values]


def test_savings_and_payback_formulas_remain_unchanged():
    tree = _app_tree()
    savings_values = _assigned_values(tree, "savings_eur")
    assert "current_price_value - total_should_cost" in [
        ast.unparse(value) for value in savings_values
    ]

    service_source = (
        ROOT / "services" / "tooling_business_case_calculation.py"
    ).read_text(encoding="utf-8")
    assert "capex / savings_per_part" in service_source
    assert "capex / annual_savings" in service_source


def test_existing_currency_context_key_is_preserved_with_new_keys():
    tree = _app_tree()
    context = next(
        node.value
        for node in ast.walk(tree)
        if isinstance(node, ast.Assign)
        and any(
            isinstance(target, ast.Subscript)
            and ast.unparse(target) == "st.session_state['part_context']"
            for target in node.targets
        )
    )
    assert isinstance(context, ast.Dict)
    keys = {
        key.value
        for key in context.keys
        if isinstance(key, ast.Constant) and isinstance(key.value, str)
    }
    assert "currency" in keys
    assert "calculation_currency" in keys
    assert "local_region_currency" in keys
